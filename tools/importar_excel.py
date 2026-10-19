"""Importa una planilla tipo 'CEDEARs en USD' a una cartera de data/db.json.

Uso:  py tools/importar_excel.py <archivo.xlsx> <id_cartera> ["Nombre cartera"]
  ej: py tools/importar_excel.py "../Mi cartera CEDEARs.xlsx" trading
      py tools/importar_excel.py "../Mi cartera dividendos.xlsx" dividendos

Reemplaza solo las operaciones, dividendos y aportes de esa cartera; el resto queda igual.
Antes de escribir guarda una copia en data/db.antes-de-importar.json.
"""
import datetime
import json
import os
import shutil
import sys
import uuid

import openpyxl

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DB_PATH = os.path.join(BASE, "data", "db.json")

DEFAULT_PORTFOLIOS = [
    {"id": "trading", "name": "Trading CEDEARs", "type": "cedear", "broker": "Cocos", "slot": 1},
    {"id": "dividendos", "name": "Dividendos", "type": "cedear", "broker": "", "slot": 2},
    {"id": "cripto", "name": "Cripto Spot", "type": "crypto", "broker": "Pionex", "slot": 3},
]


def uid():
    return uuid.uuid4().hex[:10]


def iso(d):
    if isinstance(d, datetime.datetime):
        return d.date().isoformat()
    return None


def is_num(v):
    return isinstance(v, (int, float))


def read_sheet(xlsx, pid):
    ws = openpyxl.load_workbook(xlsx, data_only=True)["CEDEARs"]
    txs, deposits, dividends = [], [], []

    def buy_from_row(r):
        broker = ws.cell(r, 2).value
        return {
            "id": uid(), "portfolioId": pid, "side": "buy",
            "ticker": str(ws.cell(r, 1).value).strip().upper(),
            "date": iso(ws.cell(r, 6).value),
            "qty": float(ws.cell(r, 3).value),
            "ratio": float(ws.cell(r, 4).value),
            "price": float(ws.cell(r, 8).value), "currency": "ARS",
            "fx": float(ws.cell(r, 7).value), "fee": 0,
            "note": "Importado del Excel" + (f" · {broker}" if broker else ""),
        }

    def add_dividend(r, t):
        m = ws.cell(r, 13).value
        if is_num(m) and m:
            dividends.append({"id": uid(), "portfolioId": pid, "ticker": t["ticker"], "date": t["date"],
                              "amount": float(m), "note": "Importado del Excel (fecha aproximada)"})

    # Operaciones abiertas: filas 4..45
    for r in range(4, 46):
        if ws.cell(r, 1).value and is_num(ws.cell(r, 3).value):
            t = buy_from_row(r)
            txs.append(t)
            add_dividend(r, t)

    # Operaciones cerradas: desde la fila 52 hasta la fila TOTAL
    r = 52
    while ws.cell(r, 1).value != "TOTAL" and r < 200:
        if ws.cell(r, 1).value and is_num(ws.cell(r, 3).value):
            b = buy_from_row(r)
            txs.append(b)
            add_dividend(r, b)
            txs.append({
                "id": uid(), "portfolioId": pid, "side": "sell",
                "ticker": b["ticker"], "date": iso(ws.cell(r, 11).value),
                "qty": b["qty"], "ratio": b["ratio"],
                "price": float(ws.cell(r, 12).value) / b["ratio"],  # precio subyacente USD -> por CEDEAR
                "currency": "USD", "fx": None, "fee": 0, "note": "Importado del Excel",
            })
        r += 1

    # Aportes: tabla "Broker | Fecha | Monto" a la derecha (su posición varía entre planillas)
    col = next((c for c in range(19, 45) for hr in (2, 3)
                if ws.cell(hr, c).value == "Broker" and ws.cell(hr, c + 1).value == "Fecha compra"), None)
    if col:
        for r in range(3, 40):
            broker, fecha, monto = ws.cell(r, col).value, ws.cell(r, col + 1).value, ws.cell(r, col + 2).value
            extra = ws.cell(r, col + 3).value
            if broker and is_num(monto):
                d = iso(fecha)
                note = str(broker)
                if not d:
                    note += f" - {int(fecha) if is_num(fecha) else fecha}"
                    d = f"{int(fecha)}-12-31" if is_num(fecha) else "2025-12-31"
                if isinstance(extra, str) and extra.strip():
                    note += f" ({extra.strip()})"
                deposits.append({"id": uid(), "portfolioId": pid, "date": d, "amount": float(monto), "note": note})
    return txs, dividends, deposits


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(1)
    xlsx, pid = sys.argv[1], sys.argv[2]
    name = sys.argv[3] if len(sys.argv) > 3 else None

    if os.path.exists(DB_PATH):
        with open(DB_PATH, encoding="utf-8") as f:
            db = json.load(f)
        shutil.copy(DB_PATH, DB_PATH.replace("db.json", "db.antes-de-importar.json"))
    else:
        db = {"version": 1, "portfolios": [dict(p) for p in DEFAULT_PORTFOLIOS], "transactions": [],
              "dividends": [], "deposits": [], "manualPrices": {}, "settings": {"baseCurrency": "USD"}}

    if not any(p["id"] == pid for p in db["portfolios"]):
        used = {p.get("slot") for p in db["portfolios"]}
        slot = next(s for s in range(1, 9) if s not in used)
        db["portfolios"].append({"id": pid, "name": name or pid, "type": "cedear", "broker": "", "slot": slot})

    txs, dividends, deposits = read_sheet(xlsx, pid)
    p = next(p for p in db["portfolios"] if p["id"] == pid)
    if not p.get("broker"):
        brokers = sorted({t["note"].split(" · ")[1] for t in txs if " · " in t["note"]})
        p["broker"] = ", ".join(brokers)
    keep = lambda rows: [x for x in rows if x["portfolioId"] != pid]
    db["transactions"] = sorted(keep(db["transactions"]) + txs, key=lambda t: t["date"])
    db["dividends"] = keep(db.get("dividends", [])) + dividends
    db["deposits"] = keep(db.get("deposits", [])) + deposits

    with open(DB_PATH, "w", encoding="utf-8") as f:
        json.dump(db, f, ensure_ascii=False, indent=2)
    print(f"Cartera '{pid}': {len(txs)} operaciones, {len(dividends)} dividendos, {len(deposits)} aportes")


if __name__ == "__main__":
    main()
